"""ANA-4 -- spreadsheet exports with validated columns and full auditing.

Per LLD section 12 an export of 5000 rows or fewer is built synchronously and
streamed; anything larger defers ``build_export`` and is polled.  Both paths
run the *same* row builder, so the file a coordinator downloads immediately and
the file the worker produces an hour later are the same file.

Three shapes the design is deliberate about:

* **The command writes a request; the route streams bytes.**  Building a
  workbook is I/O, and `decide` is pure -- so `request_export` records what was
  asked for, audits it, and decides which path it takes, while
  ``GET /api/v1/exports/{id}/download`` does the building.  The alternative,
  generating bytes inside a command, would put openpyxl inside the write
  transaction for no benefit.
* **Columns are refused, not tolerated.**  A column outside the registry is a
  rejection with every offending name listed at once, rather than a blank
  spreadsheet column nobody notices until a company asks what it means.
* **Audit metadata stays in the portal.** The audit log records the requester,
  columns, and filters; CSV/XLSX files begin directly with the table header so
  downstream spreadsheet tools can consume them without skipping preamble rows.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Literal, cast
from uuid import UUID, uuid4

import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict
from sqlalchemy.ext.asyncio import AsyncConnection, AsyncSession

from app.core.errors import (
    CYCLE_NOT_FOUND,
    EXPORT_NOT_FOUND,
    JOB_NOT_FOUND,
    UNKNOWN_EXPORT_COLUMN,
)
from app.core.plan import (
    ActorContext,
    Deferred,
    Plan,
    Reason,
    Rejection,
    ScopeIds,
    StateOp,
)
from app.core.registry import Registry
from app.modules.analytics.columns import (
    QUESTION_PREFIX,
    ExportColumn,
    application_registry,
    default_columns,
    membership_registry,
    unknown_columns,
)
from app.modules.offers.derivations import (
    JOB_APPLICATION_HAS_OFFER_SQL,
    JOB_APPLICATION_OFFER_EXPORT_LATERAL_SQL,
)

#: LLD section 12: at or below this many rows an export is built synchronously
#: and streamed; above it the request defers build_export and is polled.
SYNC_ROW_LIMIT = 5000

ExportKind = Literal["job_applications", "cycle_memberships"]
ExportFormat = Literal["xlsx", "csv"]


class RequestExportInput(BaseModel):
    model_config = ConfigDict(extra="forbid")

    cycle_id: UUID
    kind: ExportKind
    job_id: UUID | None = None
    columns: list[str] = []
    format: ExportFormat = "xlsx"
    #: Per-job exports may be narrowed to one reached round (ANA-4).
    round_id: UUID | None = None
    status: str | None = None
    cohort: Literal["applicants", "round", "offers"] = "applicants"
    #: A round export normally contains only students who passed that round.
    passed_only: bool = True


class ExportSummary(BaseModel):
    export_id: UUID
    kind: str
    format: str
    columns: list[str]
    row_count: int
    mode: str
    status: str


@dataclass(frozen=True, slots=True)
class ExportState:
    scope_ids: ScopeIds
    cycle_archived: bool
    cycle_exists: bool
    job_exists: bool
    registry: tuple[ExportColumn, ...]
    row_count: int
    export_id: UUID


def _registry_for(kind: str, questions: tuple[tuple[UUID, str], ...]) -> tuple[ExportColumn, ...]:
    if kind == "cycle_memberships":
        return membership_registry()
    return application_registry(questions)


async def _load_export(tx: AsyncSession, input_value: BaseModel, *, lock: bool) -> ExportState:
    if not isinstance(input_value, RequestExportInput):
        raise TypeError("request_export requires RequestExportInput")
    cycle = (
        (
            await tx.execute(
                sa.text("SELECT archived_at FROM cycles WHERE id = :id"),
                {"id": input_value.cycle_id},
            )
        )
        .mappings()
        .one_or_none()
    )
    job_exists = True
    questions: tuple[tuple[UUID, str], ...] = ()
    if input_value.kind == "job_applications":
        job = (
            (
                await tx.execute(
                    sa.text("SELECT id FROM jobs WHERE id = :id AND cycle_id = :cycle"),
                    {"id": input_value.job_id, "cycle": input_value.cycle_id},
                )
            )
            .mappings()
            .one_or_none()
        )
        job_exists = job is not None
        if job is not None:
            rows = (
                (
                    await tx.execute(
                        sa.text(
                            "SELECT id, text FROM job_questions WHERE job_id = :job_id ORDER BY ord"
                        ),
                        {"job_id": input_value.job_id},
                    )
                )
                .mappings()
                .all()
            )
            questions = tuple((cast(UUID, row["id"]), str(row["text"])) for row in rows)
    count_sql, count_params = _row_count_query(input_value)
    row_count = (
        int(await tx.scalar(sa.text(count_sql), count_params) or 0)
        if cycle is not None and job_exists
        else 0
    )
    return ExportState(
        scope_ids=ScopeIds(cycle_id=input_value.cycle_id, job_id=input_value.job_id),
        cycle_archived=cycle is not None and cycle["archived_at"] is not None,
        cycle_exists=cycle is not None,
        job_exists=job_exists,
        registry=_registry_for(input_value.kind, questions),
        row_count=row_count,
        export_id=uuid4(),
    )


def _row_count_query(
    input_value: RequestExportInput,
) -> tuple[str, dict[str, object]]:
    if input_value.kind == "cycle_memberships":
        return (
            "SELECT count(*) FROM cycle_memberships WHERE cycle_id = :cycle_id",
            {"cycle_id": input_value.cycle_id},
        )
    return (
        "SELECT count(*) FROM applications a WHERE a.job_id = :job_id "
        "AND (CAST(:status AS text) IS NULL "
        "     OR CAST(a.status AS text) = CAST(:status AS text)) "
        "AND (CAST(:round_id AS uuid) IS NULL "
        "     OR EXISTS (SELECT 1 FROM application_round_states rs "
        "                WHERE rs.application_id = a.id "
        "                  AND rs.round_id = CAST(:round_id AS uuid))) "
        "AND (CAST(:round_id AS uuid) IS NULL OR NOT :passed_only "
        "     OR EXISTS (SELECT 1 FROM application_round_states rs "
        "                WHERE rs.application_id = a.id "
        "                  AND rs.round_id = CAST(:round_id AS uuid) "
        "                  AND rs.result = 'advanced')) "
        "AND (CAST(:cohort AS text) <> 'offers' OR "
        f"     {JOB_APPLICATION_HAS_OFFER_SQL})",
        {
            "job_id": input_value.job_id,
            "status": input_value.status,
            "round_id": input_value.round_id,
            "cohort": input_value.cohort,
            "passed_only": input_value.passed_only,
        },
    )


def _decide_request_export(
    input_value: BaseModel,
    state: object,
    _policy: object,
    _overrides: object,
    actor: ActorContext,
) -> Plan | Rejection:
    """Record and authorise a spreadsheet export (Behavior ANA-4)."""
    if not isinstance(input_value, RequestExportInput) or not isinstance(state, ExportState):
        raise TypeError("Invalid request_export decision input")
    if not state.cycle_exists:
        return Rejection(reasons=[Reason(code=CYCLE_NOT_FOUND, human="The cycle does not exist")])
    if input_value.kind == "job_applications" and (
        input_value.job_id is None or not state.job_exists
    ):
        return Rejection(
            reasons=[
                Reason(
                    code=JOB_NOT_FOUND,
                    human="The job does not exist in this cycle",
                    path="job_id",
                )
            ]
        )

    columns = tuple(input_value.columns) or default_columns(state.registry)
    unknown = unknown_columns(columns, state.registry)
    if unknown:
        # Every offending column at once: a picker that reports one bad name
        # per attempt is a picker somebody fights three times.
        return Rejection(
            reasons=[
                Reason(
                    code=UNKNOWN_EXPORT_COLUMN,
                    human=f"{name} is not a column of this export",
                    path=f"columns.{index}",
                )
                for index, name in enumerate(unknown)
            ]
        )

    deferred_build = state.row_count > SYNC_ROW_LIMIT
    params: dict[str, object] = {
        "cycle_id": str(input_value.cycle_id),
        "job_id": None if input_value.job_id is None else str(input_value.job_id),
        "kind": input_value.kind,
        "format": input_value.format,
        "columns": list(columns),
        "round_id": None if input_value.round_id is None else str(input_value.round_id),
        "status": input_value.status,
        "cohort": input_value.cohort,
        "passed_only": input_value.passed_only,
    }
    status = "queued" if deferred_build else "ready"
    return Plan(
        state_ops=[
            StateOp(
                op="insert",
                model="export_jobs",
                values={
                    "id": state.export_id,
                    "kind": input_value.kind,
                    "params": params,
                    "status": status,
                    "requested_by": actor.user_id,
                },
            )
        ],
        events=[],
        deferred=(
            [
                # LLD section 12: only the large path defers.  The worker
                # re-reads the request and runs the same builder the download
                # route runs.
                Deferred(task="build_export", args={"export_id": str(state.export_id)})
            ]
            if deferred_build
            else []
        ),
        audit={
            "subject_type": "export",
            "subject_id": state.export_id,
            # ANA-4: who, what, columns, filters -- the whole point of auditing
            # an export is being able to answer "who sent this out, and what
            # was in it" months later.
            "details": {
                "kind": input_value.kind,
                "format": input_value.format,
                "columns": list(columns),
                "filters": {
                    "cycle_id": str(input_value.cycle_id),
                    "job_id": (None if input_value.job_id is None else str(input_value.job_id)),
                    "round_id": (
                        None if input_value.round_id is None else str(input_value.round_id)
                    ),
                    "status": input_value.status,
                    "passed_only": input_value.passed_only,
                },
                "row_count": state.row_count,
                "mode": "deferred" if deferred_build else "sync",
            },
        },
        summary={
            "export_id": str(state.export_id),
            "kind": input_value.kind,
            "format": input_value.format,
            "columns": list(columns),
            "row_count": state.row_count,
            "mode": "deferred" if deferred_build else "sync",
            "status": status,
        },
    )


# --------------------------------------------------------------------------
# The row builder, shared by the streaming route and the worker.
# --------------------------------------------------------------------------

_APPLICATION_ROWS_SQL = f"""
    SELECT
        a.id AS application_id,
        u.email AS institute_email,
        u.full_name,
        e.roll_number,
        CASE WHEN CAST(:round_id AS uuid) IS NULL
             THEN CAST(a.status AS text)
             ELSE CAST(selected_state.result AS text)
        END AS status,
        CAST(a.status AS text) AS overall_status,
        jr.name AS current_round,
        a.applied_at,
        a.resume_url AS resume_link,
        selected_state.attendance AS selected_round_attendance,
        selected_state.result AS selected_round_result,
        coalesce(selected_state.venue_override, selected_round.venue)
          AS selected_round_venue,
        coalesce(selected_state.scheduled_at_override, selected_round.scheduled_at)
          AS selected_round_time,
        coalesce(offer_data.offer_count, 0) AS offer_count,
        offer_data.offer_statuses,
        offer_data.extended_at AS offer_extended_at,
        offer_data.deadline_at AS offer_deadline_at,
        offer_data.responded_at AS offer_responded_at,
        offer_data.termination_reason AS offer_termination_reason,
        p.program_id, p.primary_branch_id, p.secondary_branch_id,
        p.graduating_year, p.cpi, p.active_backlogs, p.total_backlogs,
        CAST(p.gender AS text) AS gender, p.personal_email, p.contact_number,
        p.nationality, p.tenth_percent, p.tenth_year, p.twelfth_percent,
        p.twelfth_year, p.minor1_id, p.minor2_id, p.github_url, p.linkedin_url,
        p.portfolio_url
    FROM applications a
    JOIN enrollments e ON e.id = a.enrollment_id
    JOIN users u ON u.id = e.user_id
    LEFT JOIN profiles p ON p.enrollment_id = a.enrollment_id
    LEFT JOIN job_rounds jr ON jr.id = a.current_round_id
    LEFT JOIN application_round_states selected_state
      ON selected_state.application_id = a.id
     AND selected_state.round_id = CAST(:round_id AS uuid)
    LEFT JOIN job_rounds selected_round ON selected_round.id = selected_state.round_id
    LEFT JOIN LATERAL (
        {JOB_APPLICATION_OFFER_EXPORT_LATERAL_SQL}
    ) offer_data ON true
    WHERE a.job_id = :job_id
      AND (CAST(:status AS text) IS NULL
           OR CAST(a.status AS text) = CAST(:status AS text))
      AND (CAST(:round_id AS uuid) IS NULL
           OR EXISTS (SELECT 1 FROM application_round_states rs
                      WHERE rs.application_id = a.id
                        AND rs.round_id = CAST(:round_id AS uuid)))
      AND (CAST(:round_id AS uuid) IS NULL OR NOT :passed_only
           OR selected_state.result = 'advanced')
      AND (CAST(:cohort AS text) <> 'offers'
           OR {JOB_APPLICATION_HAS_OFFER_SQL})
    ORDER BY u.full_name, a.id
"""

_MEMBERSHIP_ROWS_SQL = """
    SELECT
        u.email AS institute_email,
        u.full_name,
        e.roll_number,
        CAST(m.status AS text) AS membership_status,
        CAST(m.outcome_tag AS text) AS outcome_tag,
        m.auto_created,
        m.consented_at,
        m.decided_at,
        m.rejection_reason,
        p.program_id, p.primary_branch_id, p.secondary_branch_id,
        p.graduating_year, p.cpi, p.active_backlogs, p.total_backlogs,
        CAST(p.gender AS text) AS gender, p.personal_email, p.contact_number,
        p.nationality, p.tenth_percent, p.tenth_year, p.twelfth_percent,
        p.twelfth_year, p.minor1_id, p.minor2_id, p.github_url, p.linkedin_url,
        p.portfolio_url
    FROM cycle_memberships m
    JOIN enrollments e ON e.id = m.enrollment_id
    JOIN users u ON u.id = e.user_id
    LEFT JOIN profiles p ON p.enrollment_id = m.enrollment_id
    WHERE m.cycle_id = :cycle_id
    ORDER BY u.full_name, m.id
"""

_ANSWERS_SQL = """
    SELECT aa.application_id, aa.question_id, aa.value
    FROM application_answers aa
    JOIN applications a ON a.id = aa.application_id
    WHERE a.job_id = :job_id
"""

#: Taxonomy columns hold ids; a spreadsheet wants the name a human reads.
_TAXONOMY_COLUMNS = {
    "program_id": "programs",
    "primary_branch_id": "branches",
    "secondary_branch_id": "branches",
    "minor1_id": "minors",
    "minor2_id": "minors",
}


async def _taxonomy_names(connection: AsyncConnection) -> dict[UUID, str]:
    names: dict[UUID, str] = {}
    for table in ("programs", "branches", "minors"):
        rows = (await connection.execute(sa.text(f"SELECT id, name FROM {table}"))).mappings().all()
        for row in rows:
            names[cast(UUID, row["id"])] = str(row["name"])
    return names


async def build_table(connection: AsyncConnection, params: dict[str, object]) -> list[list[object]]:
    """Header plus body for one export request.

    The single builder both delivery paths use.  If the streamed file and the
    deferred file could be produced by different code, they would eventually
    say different things about the same students.
    """
    columns = [str(column) for column in cast("list[object]", params["columns"])]
    kind = str(params["kind"])
    if kind == "cycle_memberships":
        rows = (
            (
                await connection.execute(
                    sa.text(_MEMBERSHIP_ROWS_SQL), {"cycle_id": params["cycle_id"]}
                )
            )
            .mappings()
            .all()
        )
        answers: dict[UUID, dict[UUID, object]] = {}
        labels = {column.key: column.label for column in membership_registry()}
    else:
        rows = (
            (
                await connection.execute(
                    sa.text(_APPLICATION_ROWS_SQL),
                    {
                        "job_id": params["job_id"],
                        "status": params.get("status"),
                        "round_id": params.get("round_id"),
                        "cohort": params.get("cohort", "applicants"),
                        "passed_only": params.get("passed_only", True),
                    },
                )
            )
            .mappings()
            .all()
        )
        answer_rows = (
            (await connection.execute(sa.text(_ANSWERS_SQL), {"job_id": params["job_id"]}))
            .mappings()
            .all()
        )
        answers = {}
        for row in answer_rows:
            answers.setdefault(cast(UUID, row["application_id"]), {})[
                cast(UUID, row["question_id"])
            ] = row["value"]
        # The questions have to be loaded for their *labels*, not only for
        # validation: without them a question column exports under its raw
        # `question_<uuid>` key, which is unreadable in the one artefact that
        # leaves this system.
        question_rows = (
            (
                await connection.execute(
                    sa.text(
                        "SELECT id, text FROM job_questions WHERE job_id = :job_id ORDER BY ord"
                    ),
                    {"job_id": params["job_id"]},
                )
            )
            .mappings()
            .all()
        )
        questions = tuple((cast(UUID, row["id"]), str(row["text"])) for row in question_rows)
        labels = {column.key: column.label for column in application_registry(questions)}

    names = await _taxonomy_names(connection)
    header: list[object] = [labels.get(column, column) for column in columns]
    table: list[list[object]] = [header]
    for row in rows:
        mapping = dict(row)
        application_id = cast("UUID | None", mapping.get("application_id"))
        line: list[object] = []
        for column in columns:
            if column.startswith(QUESTION_PREFIX):
                question_id = UUID(column[len(QUESTION_PREFIX) :])
                value = answers.get(application_id or UUID(int=0), {}).get(question_id)
                line.append(_cell(value, names))
                continue
            line.append(_cell(mapping.get(column), names))
        table.append(line)
    return table


def _cell(value: object, names: dict[UUID, str]) -> object:
    """One value as a spreadsheet cell: names for ids, ISO for instants."""
    if value is None:
        return ""
    if isinstance(value, UUID):
        return names.get(value, str(value))
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, list):
        return ", ".join(str(item) for item in cast("list[object]", value))
    return value


def to_csv(table: list[list[object]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerows(table)
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(table: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:  # pragma: no cover - openpyxl always creates one
        sheet = workbook.create_sheet()
    sheet.title = "Export"
    for line in table:
        sheet.append(line)

    header_row = 1
    sheet.freeze_panes = f"A{header_row + 1}"
    last_column = get_column_letter(max(1, len(table[0])))
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{sheet.max_row}"
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[header_row]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column_index in range(1, max(1, len(table[0])) + 1):
        values = [
            str(sheet.cell(row=row, column=column_index).value or "")
            for row in range(header_row, sheet.max_row + 1)
        ]
        width = min(55, max(12, max(map(len, values), default=12) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width
        for row_index in range(header_row + 1, sheet.max_row + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render(
    table: list[list[object]],
    params: dict[str, object],
    *,
    requested_by: str,
    built_at: datetime | None = None,
) -> tuple[bytes, str, str]:
    """Bytes, media type, and filename for one built export."""
    moment = built_at or datetime.now(UTC)
    stamp = moment.strftime("%Y%m%d-%H%M")
    if str(params["format"]) == "csv":
        return (
            to_csv(table),
            "text/csv; charset=utf-8",
            f"{params['kind']}-{stamp}.csv",
        )
    return (
        to_xlsx(table),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        f"{params['kind']}-{stamp}.xlsx",
    )


def register_export_commands(registry: Registry) -> None:
    registry.command(
        name="request_export",
        input_model=RequestExportInput,
        output_model=ExportSummary,
        actor="staff",
        scope="cycle",
        loader=_load_export,
        rule_domains=(),
        spec_ids=("ANA-4",),
        rate_limit="5/min",
    )(_decide_request_export)
    registry.command(
        name="complete_export",
        input_model=CompleteExportInput,
        output_model=CompleteExportSummary,
        actor="system",
        scope="none",
        loader=_load_complete_export,
        rule_domains=(),
        spec_ids=("ANA-4",),
        expose_http=False,
    )(_decide_complete_export)


# --------------------------------------------------------------------------
# The deferred path: build_export's completion, as a command like everything
# else that writes.
# --------------------------------------------------------------------------

#: LLD section 12: a built result is held in export_jobs.result_meta as base64
#: up to 10 MB; anything larger is not stored and the download rebuilds it.
MAX_STORED_RESULT_BYTES = 10 * 1024 * 1024


class CompleteExportInput(BaseModel):
    model_config = ConfigDict(extra="forbid")

    export_id: UUID
    encoded: str | None = None
    byte_size: int = 0
    error: str | None = None


class CompleteExportSummary(BaseModel):
    export_id: UUID
    status: str
    stored: bool


@dataclass(frozen=True, slots=True)
class CompleteExportState:
    scope_ids: ScopeIds
    exists: bool
    already_ready: bool


async def _load_complete_export(
    tx: AsyncSession, input_value: BaseModel, *, lock: bool
) -> CompleteExportState:
    if not isinstance(input_value, CompleteExportInput):
        raise TypeError("complete_export requires CompleteExportInput")
    row = (
        (
            await tx.execute(
                sa.text(
                    "SELECT status FROM export_jobs WHERE id = :id"
                    + (" FOR UPDATE" if lock else "")
                ),
                {"id": input_value.export_id},
            )
        )
        .mappings()
        .one_or_none()
    )
    return CompleteExportState(
        scope_ids=ScopeIds(),
        exists=row is not None,
        already_ready=row is not None and str(row["status"]) == "ready",
    )


def _decide_complete_export(
    input_value: BaseModel,
    state: object,
    _policy: object,
    _overrides: object,
    _actor: ActorContext,
) -> Plan | Rejection:
    """Record what the build_export worker produced (LLD section 12).

    Idempotent by design: the task may run twice, and a second completion of an
    export already marked ready is a no-op rather than a second write, because
    a deferred handler that is not idempotent is a handler that corrupts state
    on any retry.
    """
    if not isinstance(input_value, CompleteExportInput) or not isinstance(
        state, CompleteExportState
    ):
        raise TypeError("Invalid complete_export decision input")
    if not state.exists:
        return Rejection(reasons=[Reason(code=EXPORT_NOT_FOUND, human="No such export")])
    if state.already_ready:
        return Plan(
            state_ops=[],
            events=[],
            deferred=[],
            audit=None,
            summary={
                "export_id": str(input_value.export_id),
                "status": "ready",
                "stored": False,
            },
        )

    failed = input_value.error is not None
    stored = (
        not failed
        and input_value.encoded is not None
        and input_value.byte_size <= MAX_STORED_RESULT_BYTES
    )
    result_meta: dict[str, object] = {
        "byte_size": input_value.byte_size,
        "stored": stored,
        # An oversize result is not kept; the download route rebuilds it, which
        # is what LLD section 12 means by "else re-run".
        "content": input_value.encoded if stored else None,
    }
    return Plan(
        state_ops=[
            StateOp(
                op="update",
                model="export_jobs",
                values={
                    "status": "failed" if failed else "ready",
                    "result_meta": None if failed else result_meta,
                    "error": input_value.error,
                },
                where={"id": input_value.export_id},
            )
        ],
        events=[],
        deferred=[],
        audit=None,
        summary={
            "export_id": str(input_value.export_id),
            "status": "failed" if failed else "ready",
            "stored": stored,
        },
    )
